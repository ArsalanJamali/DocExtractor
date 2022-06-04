from django.http import HttpResponse
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from knox.auth import TokenAuthentication
from .serializers import DocumentSerializer,DocumentUpdateSerializer,DocumentImageSerializer
from .models import Document, DocumentImage,Label
from rest_framework.decorators import api_view,authentication_classes,permission_classes
import uuid
from .utils import process_image
from PIL import Image
import json
from openpyxl import Workbook
from rest_framework.generics import UpdateAPIView,DestroyAPIView,ListAPIView
from django.db.models import Q 

# Create your views here.





@api_view(['GET',])
@authentication_classes([TokenAuthentication,])
@permission_classes([IsAuthenticated])
def get_recent_models(request):
    documents=Document.objects.filter(user=request.user.pk).order_by('-created_at','-pk')
    return Response(DocumentSerializer(documents,many=True).data)

@api_view(['GET','POST'])
@authentication_classes([TokenAuthentication,])
@permission_classes([IsAuthenticated])
def create_document_model(request):
    if request.method=='GET':
        return Response({
            'model_id':uuid.uuid4(),
        })
    else:
        model_id=request.POST['model_id']
        type=request.POST['type']
        description=request.POST['description']
        try:
            doc=Document.objects.get(model_id=model_id)
        except:
            doc=None
        
        if doc:
            return Response({'error':'Model with this id already exist'})

        doc=Document(user=request.user,type=type,model_id=model_id,description=description)
        doc.save()
        return Response(DocumentSerializer(doc).data)

@api_view(['POST'])
@authentication_classes([TokenAuthentication,])
@permission_classes([IsAuthenticated])
def process_documents(request):
    images=request.FILES.getlist('files')
    pk=request.POST['pk']
    
    try:
        doc=Document.objects.get(pk=pk)
    except:
        doc=None
    
    if doc==None:
        return Response({'error':'No such model exists!'})
    
    returned_dict=dict()

    for image in images:
        doc_image_obj=DocumentImage(document=doc,document_image=image)
        image=Image.open(image)
        output=process_image(image,doc.type)
        doc_image_obj.save()
        url=doc_image_obj.document_image.url
        for key,value in output.items():
            obj=Label(image=doc_image_obj,key=key,value=value)
            obj.save()

        returned_dict[doc_image_obj.pk]={
            'url':url,
            'output':output
        }

    return Response(returned_dict)

@api_view(['POST'])
@authentication_classes([TokenAuthentication,])
@permission_classes([IsAuthenticated])
def save_json(request):
    pk_list=request.POST.getlist('pk')
    error=False
    objs=DocumentImage.objects.filter(pk__in=pk_list)
    
    if(len(objs)!=len(pk_list)):
        error=True

    output_dict=dict()
    if error:
        output_dict[error]="Some file data is missing!"
    
    for obj in objs:
        filename=obj.document_image.name.split('/')[1]
        output_dict[filename]=dict()
        for label_obj in obj.label_set.all():
            output_dict[filename][label_obj.key]=label_obj.value
    
        

    response = HttpResponse(json.dumps(output_dict),content_type="application/json")
    response['Content-Disposition'] = 'attachment; filename=data.json'
    return response  

@api_view(['POST'])
@authentication_classes([TokenAuthentication,])
@permission_classes([IsAuthenticated])
def save_csv(request):
    pk_list=request.POST.getlist('pk')
    objs=DocumentImage.objects.filter(pk__in=pk_list)
    
    csv=Workbook()
    flag=True
    sheet=None
    for obj in objs:
        if flag:
            sheet=csv.active
            flag=False
        else:
            sheet=csv.create_sheet()
        sheet.title=filename=obj.document_image.name.split('/')[1]
        for label_obj in obj.label_set.all():
            sheet.append([label_obj.key,label_obj.value])
            
    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response['Content-Disposition'] = 'attachment; filename=data.xlsx'
    
    csv.save(response)
    return response

class UpdateModelDescription(UpdateAPIView):
    serializer_class=DocumentUpdateSerializer
    authentication_classes=(TokenAuthentication,)
    permission_classes=(IsAuthenticated,)
    
    def get_queryset(self):
        return  Document.objects.filter(user=self.request.user.pk)

class DeleteModelView(DestroyAPIView):
    authentication_classes=(TokenAuthentication,)
    permission_classes=(IsAuthenticated,)
    
    def get_queryset(self):
        return  Document.objects.filter(user=self.request.user.pk)

class DeleteModelImageView(DestroyAPIView):
    authentication_classes=(TokenAuthentication,)
    permission_classes=(IsAuthenticated,)
    
    def get_queryset(self):
        return  DocumentImage.objects.all()

class ListModelImages(ListAPIView):
    serializer_class=DocumentImageSerializer
    authentication_classes=(TokenAuthentication,)
    permission_classes=(IsAuthenticated,)
    
    def get_queryset(self):
        return  DocumentImage.objects.filter(document=self.kwargs.get('pk'))

@api_view(['POST'])
@authentication_classes([TokenAuthentication,])
@permission_classes([IsAuthenticated])
def preview_all_images(request):
    pk_list=request.POST.getlist('pk')
    objs=DocumentImage.objects.filter(pk__in=pk_list)
    output_list=list()
    
    for obj in objs:
        output_dict=dict()
        output_dict['url']='' if obj.document_image==None else obj.document_image.url
        output_dict['label']=dict()
        output_dict['image_id']=obj.pk
        for label_obj in obj.label_set.all():
            output_dict['label'][label_obj.key]=label_obj.value

        output_list.append(output_dict)

    return Response(output_list)

@api_view(['DELETE'])
@authentication_classes([TokenAuthentication,])
@permission_classes([IsAuthenticated])
def delete_all_images(request):
    pk_list=request.POST.getlist('pk')
    DocumentImage.objects.filter(pk__in=pk_list).delete()
    return Response({'Success':True})

@api_view(['GET',])
@authentication_classes([TokenAuthentication,])
@permission_classes([IsAuthenticated])
def SearchModel(request):
    query=request.GET['query']
    
    if query=='':
        return Response("Empty Query")


    
    documents=Document.objects.filter( Q(user=request.user.pk) & (Q(model_id__icontains=query) | Q(description__icontains=query)) ).order_by('-created_at','-pk')
    
    return Response(DocumentSerializer(documents,many=True).data)



@api_view(['DELETE'])
@authentication_classes([TokenAuthentication,])
@permission_classes([IsAuthenticated])
def delete_all_images(request):
    pk_list=request.POST.getlist('pk')
    DocumentImage.objects.filter(pk__in=pk_list).delete()
    return Response({'Success':True})

# apply try except remember


#    threads=list()
# for image in images:
#         image=Image.open(image)
#         process_image(image)
        # process = Thread(target=process_image, args=[image])
        # process.start()
        # threads.append(process)
    
    # for process in threads:
    #     process.join()


